from pathlib import Path
import sqlite3

import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

GREEN_FILL = PatternFill(
    fill_type="solid",
    start_color="92D050",
    end_color="92D050",
)

YELLOW_FILL = PatternFill(
    fill_type="solid",
    start_color="FFD966",
    end_color="FFD966",
)

RED_FILL = PatternFill(
    fill_type="solid",
    start_color="F4CCCC",
    end_color="F4CCCC",
)

BENCHMARK_FILL = PatternFill(
    fill_type="solid",
    start_color="FFC000",
    end_color="FFC000",
)

BASE_DIR = Path(__file__).resolve().parents[2]

DATABASE = BASE_DIR / "data" / "database" / "nifty100.db"

OUTPUT_DIR = BASE_DIR / "output"

OUTPUT_DIR.mkdir(exist_ok=True)

OUTPUT_FILE = OUTPUT_DIR / "peer_comparison.xlsx"

def load_tables():

    conn = sqlite3.connect(DATABASE)

    comparison = pd.read_sql(
        "SELECT * FROM comparison_table",
        conn,
    )

    percentiles = pd.read_sql(
        "SELECT * FROM peer_percentiles",
        conn,
    )

    peer = pd.read_sql(
        "SELECT * FROM peer_groups",
        conn,
    )

    conn.close()

    return comparison, percentiles, peer

METRICS = [

    "roe_percentage",

    "roce_percentage",

    "net_profit_margin_pct",

    "operating_profit_margin_pct",

    "return_on_assets_pct",

    "return_on_equity_pct",

    "return_on_capital_employed_pct",

    "debt_to_equity",

    "interest_coverage",

    "asset_turnover",

    "free_cash_flow_cr",

    "cfo_quality_score",

    "fcf_conversion_rate",

    "revenue_cagr_5yr",

    "pat_cagr_5yr",

    "eps_cagr_5yr",

    "composite_quality_score",
]

def build_percentile_table(percentiles):

    percentile_table = (
        percentiles
        .pivot(
            index="company_id",
            columns="metric",
            values="percentile_rank",
        )
        .reset_index()
    )

    percentile_table.columns.name = None

    percentile_table = percentile_table.rename(
    columns={
        column: f"{column}_percentile"
        for column in percentile_table.columns
        if column != "company_id"
    }
)

    return percentile_table

def merge_tables(
    comparison,
    percentile_table,
    peer,
):

    merged = comparison.merge(
        percentile_table,
        on="company_id",
        how="left",
    )

    merged = merged.merge(
        peer[
            [
                "company_id",
                "is_benchmark",
            ]
        ],
        on="company_id",
        how="left",
    )

    return merged

def write_excel(merged):

    with pd.ExcelWriter(
        OUTPUT_FILE,
        engine="openpyxl",
    ) as writer:

        for group in sorted(
            merged["peer_group_name"]
            .dropna()
            .unique()
        ):
            print("Creating sheet:", group)

            sheet = (
                merged[
                    merged["peer_group_name"] == group
                ]
                .copy()
            )

            sheet.to_excel(
                writer,
                sheet_name=group[:31],
                index=False,
            )


def apply_percentile_colours(workbook):

    percentile_columns = [
    "asset_turnover_percentile",
    "debt_to_equity_percentile",
    "eps_cagr_5yr_percentile",
    "free_cash_flow_cr_percentile",
    "interest_coverage_percentile",
    "net_profit_margin_pct_percentile",
    "pat_cagr_5yr_percentile",
    "revenue_cagr_5yr_percentile",
    "roce_percentage_percentile",
    "roe_percentage_percentile",
]

    for ws in workbook.worksheets:

        headers = [
            cell.value
            for cell in ws[1]
        ]

        for metric in percentile_columns:

            if metric not in headers:
                continue

            col = headers.index(metric) + 1

            for row in range(2, ws.max_row + 1):

                value = ws.cell(
                    row=row,
                    column=col,
                ).value

                if value is None:
                    continue

                if value >= 75:

                    ws.cell(
                        row=row,
                        column=col,
                    ).fill = GREEN_FILL

                elif value <= 25:

                    ws.cell(
                        row=row,
                        column=col,
                    ).fill = RED_FILL

                else:

                    ws.cell(
                        row=row,
                        column=col,
                    ).fill = YELLOW_FILL

def highlight_benchmark_rows(workbook):

    for ws in workbook.worksheets:

        headers = [
            cell.value
            for cell in ws[1]
        ]

        if "is_benchmark" not in headers:
            continue

        benchmark_column = (
            headers.index("is_benchmark") + 1
        )

        for row in range(2, ws.max_row + 1):

            value = ws.cell(
                row=row,
                column=benchmark_column,
            ).value

            if value in (1, True):

                for col in range(
                    1,
                    ws.max_column + 1,
                ):

                    ws.cell(
                        row=row,
                        column=col,
                    ).fill = BENCHMARK_FILL

def add_median_rows(workbook):

    from statistics import median

    for ws in workbook.worksheets:

        headers = [
            cell.value
            for cell in ws[1]
        ]

        median_row = ws.max_row + 1

        ws.cell(
            row=median_row,
            column=1,
        ).value = "Median"

        for col in range(2, ws.max_column + 1):

            values = []

            for row in range(2, median_row):

                value = ws.cell(
                    row=row,
                    column=col,
                ).value

                if isinstance(
                    value,
                    (int, float),
                ):

                    values.append(value)

            if values:

                ws.cell(
                    row=median_row,
                    column=col,
                ).value = round(
                    median(values),
                    2,
                )

def highlight_benchmark(workbook):

    for ws in workbook.worksheets:

        headers = [
            cell.value
            for cell in ws[1]
        ]

        benchmark_column = headers.index(
            "is_benchmark"
        ) + 1

        for row in range(
            2,
            ws.max_row + 1,
        ):

            value = ws.cell(
                row=row,
                column=benchmark_column,
            ).value

            if value == 1:

                for col in range(
                    1,
                    ws.max_column + 1,
                ):

                    ws.cell(
                        row=row,
                        column=col,
                    ).fill = BENCHMARK_FILL


comparison, percentiles, peer = load_tables()


percentile_table = percentiles

conn = sqlite3.connect(DATABASE)

percentile_table.to_sql(
    "peer_percentiles",
    conn,
    if_exists="replace",
    index=False,
)

conn.close()



merged = merge_tables(
    comparison,
    percentile_table,
    peer,
)

print(merged["peer_group_name"].value_counts())

write_excel(
    merged,
)
workbook = load_workbook(
    OUTPUT_FILE
)

apply_percentile_colours(
    workbook,
)

highlight_benchmark_rows(
    workbook,
)

add_median_rows(
    workbook,
)

workbook.save(
    OUTPUT_FILE
)


print(OUTPUT_FILE)

ws = workbook.worksheets[0]

headers = [cell.value for cell in ws[1]]

print(headers)


